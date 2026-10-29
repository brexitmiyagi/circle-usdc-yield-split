# Builds results/circle_model.xlsx: the same three-statement model as model.py, in a spreadsheet with live formulas,
# plus the consensus table, the price target, the Hyperliquid measurement and the options check.
# Inputs are blue, formulas black, links to another sheet green. Change any blue cell and the sheet recalculates.
# Run model.py, hyperliquid_q3.py, valuation.py and options_check.py first (run_all.py does).
import os, sys, re
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from common import read, num, market, RESULTS
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.comments import Comment

m = market()
A = {r['item']: r for r in read('model_assumptions.csv')}
MI = {r['item']: r['value'] for r in read('model_inputs.csv', folder='results')}
HL = {r['item']: num(r['value']) for r in read('hyperliquid_q3.csv', folder='results')}
COE = {r['item']: num(r['value']) for r in read('cost_of_equity.csv', folder='results')}
CD = {r['item']: num(r['value']) for r in read('cost_of_equity.csv')}
DR = {r['item']: num(r['value']) for r in read('rate_drag.csv', folder='results')}
VS = {r['item']: num(r['value']) for r in read('valuation_summary.csv', folder='results')}
BT = {r['item']: num(r['value']) for r in read('backtest_summary.csv', folder='results')}
BTQ = read('backtest_keep.csv', folder='results')
VAL = {r['Scenario']: r for r in read('table_valuation.csv', folder='results')}
CONS = {r['period']: r for r in read('consensus_2026-10-28.csv')}
ACT = {r['quarter']: r for r in read('actuals_quarterly.csv')}
a = lambda k: num(A[k]['value'])
src = lambda k: A[k]['source']

F = 'Arial'
BLUE, BLACK, GREEN = Font(name=F, color='0000FF'), Font(name=F, color='000000'), Font(name=F, color='008000')
BOLD = Font(name=F, bold=True)
HEAD = Font(name=F, bold=True, color='FFFFFF')
HFILL = PatternFill('solid', fgColor='1F3A5F')
SFILL = PatternFill('solid', fgColor='E8EEF5')
YELLOW = PatternFill('solid', fgColor='FFFF00')
TOP = Border(top=Side(style='thin'))
USD = '#,##0;(#,##0);-'
USD1 = '#,##0.0;(#,##0.0);-'
BN = '#,##0.00;(#,##0.00);-'
PS = '$0.00;($0.00);-'
PCT = '0.0%;(0.0%);-'
PCT2 = '0.00%;(0.00%);-'

wb = Workbook()

# ---------------- Assumptions ----------------
ws = wb.active; ws.title = 'Assumptions'
ws['A1'] = 'Circle Internet Group (CRCL): model assumptions'; ws['A1'].font = Font(name=F, bold=True, size=13)
ws['A2'] = 'Blue = input (change it), black = formula, green = link to another sheet, yellow = key judgement. Source for every input in column C.'
ws['A2'].font = Font(name=F, italic=True)
for c, t in zip('ABC', ('Item', 'Value', 'Source / note')):
    ws[f'{c}4'] = t; ws[f'{c}4'].font = HEAD; ws[f'{c}4'].fill = HFILL
REF = {}
r = 5
def section(title):
    global r
    r += 1; ws[f'A{r}'] = title; ws[f'A{r}'].font = BOLD; ws[f'A{r}'].fill = SFILL; ws[f'B{r}'].fill = SFILL; ws[f'C{r}'].fill = SFILL; r += 1
def inp(key, label, value, source, fmt=USD1, key_judgement=False, formula=False):
    global r
    ws[f'A{r}'] = label; ws[f'B{r}'] = value; ws[f'C{r}'] = source
    ws[f'A{r}'].font = BLACK; ws[f'C{r}'].font = Font(name=F, size=9, color='555555')
    ws[f'B{r}'].font = BLACK if formula else BLUE
    ws[f'B{r}'].number_format = fmt
    if key_judgement: ws[f'B{r}'].fill = YELLOW
    REF[key] = f'Assumptions!$B${r}'; r += 1

section('USDC float ($bn)')
inp('usdc_q3', 'USDC average, Q3-26', a('usdc_avg_q3_26_bn'), src('usdc_avg_q3_26_bn'), BN)
inp('usdc_q4', 'USDC average, Q4-26 (October to date)', float(MI['usdc_avg_q4_26_bn']), 'DefiLlama daily USDC supply, 1-28 Oct 2026, calibrated to Circle\'s reported Q2 average (model.py)', BN)
inp('end_q3', 'USDC at 30 Sep 2026', float(MI['usdc_end_q3_26_bn']), 'DefiLlama, 30 Sep 2026, same calibration', BN)
inp('end_q4', 'USDC at latest date (used for 31 Dec 2026)', float(MI['usdc_end_q4_26_bn']), 'DefiLlama, latest day in data/stablecoin_supply_daily.csv, same calibration', BN)
inp('g', 'USDC growth from Q1-27, a year', float(MI['usdc_growth_jpmorgan']), "JPMorgan's $500-600bn stablecoin market by end-2028 (midpoint $550bn) from today's $302.3bn, USDC share held (model.py)", PCT, True)
inp('drag', 'Float lost per 1.00pt of reserve return above today (gained below)', DR['drag_per_100bp_pct'] / 100, 'rate_drag.py: monthly regression, all stablecoins since Jan 2023, 6 lags, bitcoin control; 95% range 12.3% to 19.0% (Drag sheet)', PCT, True)
inp('today', 'Reserve return today', a('reserve_return_today_pct') / 100, src('reserve_return_today_pct'), PCT2)

section('Reserve return')
inp('rr_q3', 'Reserve return, Q3-26', a('reserve_return_q3_26_pct') / 100, src('reserve_return_q3_26_pct'), PCT2)
inp('rr_q4', 'Reserve return, Q4-26', a('reserve_return_q4_26_pct') / 100, src('reserve_return_q4_26_pct'), PCT2, True)
inp('gap', "Circle's gap below the 3-month bill", a('bill_gap_pt') / 100, src('bill_gap_pt'), PCT2)
cd = MI['curve_date']
inp('p3', f'Treasury par yield, 3 month ({cd})', float(MI['par_3m']) / 100, 'US Treasury daily par yield curve, home.treasury.gov', PCT2)
inp('p6', f'Treasury par yield, 6 month ({cd})', float(MI['par_6m']) / 100, 'as above', PCT2)
inp('p12', f'Treasury par yield, 1 year ({cd})', float(MI['par_1y']) / 100, 'as above', PCT2)
inp('p24', f'Treasury par yield, 2 year ({cd})', float(MI['par_2y']) / 100, 'as above', PCT2)
inp('f1', 'Forward 3m bill in 3 months', f"=2*{REF['p6']}-{REF['p3']}", 'par ~ zero at the short end', PCT2, formula=True)
inp('f2', 'Forward 6m bill in 6 months', f"=2*{REF['p12']}-{REF['p6']}", '', PCT2, formula=True)
inp('f3', 'Forward 1y bill in 1 year', f"=2*{REF['p24']}-{REF['p12']}", '', PCT2, formula=True)
inp('b1', 'Bill path, Q1-27', f"=AVERAGE({REF['f1']},{REF['f2']})", '', PCT2, formula=True)
inp('b2', 'Bill path, Q2-27', f"={REF['f2']}", '', PCT2, formula=True)
inp('b3', 'Bill path, Q3-27', f"=AVERAGE({REF['f2']},{REF['f3']})", '', PCT2, formula=True)
inp('b4', 'Bill path, Q4-27', f"={REF['f3']}", '', PCT2, formula=True)

section("Circle's keep")
inp('k0', 'Keep, Q3-26 (measured, corrected for the backtest miss)', "=ROUND(Hyperliquid!$B$24,4)", 'Hyperliquid sheet: keep_method.py on the on-chain split, plus the average miss from backtest_keep.py', PCT, formula=True)
ws[REF['k0'].split('!')[1].replace('$', '')].font = GREEN
inp('y0', 'Yield the keep is measured at', a('keep_base_yield_pct') / 100, src('keep_base_yield_pct'), PCT2)
inp('ri100', 'Reserve income per +100bp ($m)', m['rate_sens_reserve_income_per_100bp_musd'], 'Circle 10-Q Q2-26, Item 3', USD)
inp('dc100', 'Distribution costs per +100bp ($m)', m['rate_sens_distribution_cost_per_100bp_musd'], 'Circle 10-Q Q2-26, Item 3', USD)
inp('marg', 'Marginal keep of a yield change (rate table)', f"=1-{REF['dc100']}/{REF['ri100']}", 'Circle\'s own rate table', PCT, formula=True)
inp('use_rt', 'Use the rate table for the keep? (1 yes, 0 no)', int(a('use_rate_table_for_keep')), 'Base 0: the backtest did worse with it (Backtest sheet)', '0', True)
inp('mu', 'Keep on a change in yield, as used', f"=IF({REF['use_rt']}=1,{REF['marg']},{REF['k0']})", '0 above means the keep stays at the Q3 level as yields move', PCT, formula=True)

section('Revenue and costs ($m a quarter unless stated)')
inp('oth_h2', 'Other revenue ex-ARC, H2-26 total', a('other_revenue_ex_arc_h2_26_musd'), src('other_revenue_ex_arc_h2_26_musd'))
inp('oth_g', 'Other revenue growth, 2027', a('other_revenue_ex_arc_growth_2027_pct') / 100, src('other_revenue_ex_arc_growth_2027_pct'), PCT)
inp('arc_q3', 'ARC token revenue, Q3-26', a('arc_revenue_q3_26_musd'), src('arc_revenue_q3_26_musd'), USD1, True)
inp('arc_q4', 'ARC token revenue, Q4-26', a('arc_revenue_q4_26_musd'), src('arc_revenue_q4_26_musd'), USD1, True)
inp('arc_27', 'ARC token revenue, 2027 total (H1)', a('arc_revenue_2027_musd'), src('arc_revenue_2027_musd'))
inp('oc', 'Other costs of revenue', a('other_costs_qtr_musd'), src('other_costs_qtr_musd'))
inp('opex26', 'Adjusted opex, FY26', a('adj_opex_fy26_musd'), src('adj_opex_fy26_musd'))
inp('opexh1', 'Adjusted opex, H1-26 actual', a('adj_opex_h1_26_musd'), src('adj_opex_h1_26_musd'))
inp('opex_g', 'Adjusted opex growth, 2027', a('adj_opex_growth_2027_pct') / 100, src('adj_opex_growth_2027_pct'), PCT)
inp('sbc', 'Stock pay, base quarter', a('sbc_qtr_musd'), src('sbc_qtr_musd'))
inp('sbc_g', 'Stock pay growth a quarter', 0.0125, 'judgement', PCT2)
inp('da', 'Depreciation and amortisation', a('da_qtr_musd'), src('da_qtr_musd'))
inp('oth_op', 'Other GAAP opex items', a('gaap_opex_other_items_qtr_musd'), src('gaap_opex_other_items_qtr_musd'))
inp('t26', 'Tax rate, 2026', a('tax_rate_2026_pct') / 100, src('tax_rate_2026_pct'), PCT)
import model as _M
inp('t27', 'Tax rate, 2027', round(_M.T27 / 100, 4), 'model.py NOL schedule: $478.8m federal NOLs at end-2025 (10-K FY25), 80% limit, 5.1% non-federal; see the tax check on the Model sheet', PCT, True)
inp('nol0', 'Federal NOLs at 31 Dec 2025 ($m)', a('federal_nol_end_2025_musd'), src('federal_nol_end_2025_musd'), USD1)
inp('pre_h1', 'Pretax income, H1-26 ($m)', a('pretax_h1_26_musd'), src('pretax_h1_26_musd'), USD1)

section('Shares and cash flow')
inp('sh0', 'Diluted shares, Q2-26 (m)', float(MI['diluted_shares_q2']), 'Circle 10-Q Q2-26', '#,##0.000')
inp('drift', 'Diluted share drift a quarter', a('share_drift_qtr_pct') / 100, src('share_drift_qtr_pct'), PCT2)
inp('bin_sh', 'Binance shares issued 17 Sep 2026 (m)', a('binance_shares_m'), src('binance_shares_m'), '#,##0.000')
inp('bin_cash', 'Binance placement cash', a('binance_placement_cash_musd'), src('binance_placement_cash_musd'))
inp('taz', 'Tazapay consideration (paid in shares)', a('tazapay_consideration_musd'), src('tazapay_consideration_musd'))
inp('price', 'CRCL price', m['crcl_price'], 'Yahoo close, 26 Oct 2026 (entry price of the call)', PS)
inp('taz_sh', 'Tazapay shares (m)', f"={REF['taz']}/{REF['price']}", 'at the CRCL price above', '#,##0.000', formula=True)
inp('capex', 'Capex', a('capex_qtr_musd'), src('capex_qtr_musd'))
inp('rsu', 'RSU tax withholding', a('rsu_withholding_qtr_musd'), src('rsu_withholding_qtr_musd'))
inp('opt', 'Option exercise proceeds', a('option_proceeds_qtr_musd'), src('option_proceeds_qtr_musd'))
inp('recv', 'ARC presale receivable collected in Q3-26', a('arc_receivable_musd'), src('arc_receivable_musd'))

section('Opening balance sheet, 30 Jun 2026 ($m)')
inp('cash0', 'Corporate cash', float(MI['cash']), 'Circle 10-Q Q2-26', USD1)
inp('seg', 'Segregated corporate stablecoins', float(MI['segregated_corporate']), 'Circle 10-Q Q2-26', USD1)
inp('excess', 'Reserves held above the USDC liability', float(MI['reserve_excess']), 'Circle 10-Q Q2-26: reserves $73,161.2m less deposits $72,927.5m', USD1)
inp('ppe0', 'Fixed assets, intangibles, goodwill', float(MI['ppe_intangibles_goodwill']), 'Circle 10-Q Q2-26: $22.2m + $446.6m + $265.7m', USD1)
inp('oa0', 'Other assets', float(MI['other_assets']), 'total assets less the lines above (includes the $20.2m ARC receivable)', USD1)
inp('def0', 'Deferred ARC presale revenue', float(MI['deferred_arc']), 'Circle 8-K 2 Jul 2026 and 10-Q Q2-26', USD1)
inp('ol0', 'Other liabilities', float(MI['other_liabilities']), 'total liabilities less deposits and deferred ARC', USD1)
inp('eq0', 'Total equity (incl. $1.4m noncontrolling interest)', float(MI['equity']), 'Circle 10-Q Q2-26 (SEC XBRL: assets less liabilities)', USD1)
ws.column_dimensions['A'].width = 52; ws.column_dimensions['B'].width = 14; ws.column_dimensions['C'].width = 110
ws.freeze_panes = 'A5'

# ---------------- Model ----------------
md = wb.create_sheet('Model')
Q = ['Q3-26', 'Q4-26', 'Q1-27', 'Q2-27', 'Q3-27', 'Q4-27']
C = ['C', 'D', 'E', 'F', 'G', 'H']
md['A1'] = 'Circle three-statement model, base case (JPMorgan path, forward-curve yields, measured Q3 keep). $m unless stated.'
md['A1'].font = Font(name=F, bold=True, size=12)
md['A3'] = 'Line'; md['B3'] = 'Unit'
for c, q in zip(C, Q): md[f'{c}3'] = q
for c in ['A', 'B'] + C: md[f'{c}3'].font = HEAD; md[f'{c}3'].fill = HFILL; md[f'{c}3'].alignment = Alignment(horizontal='center')
R = {}; mr = [4]
def hdr(t):
    mr[0] += 1; md[f'A{mr[0]}'] = t; md[f'A{mr[0]}'].font = BOLD
    for c in ['A', 'B'] + C: md[f'{c}{mr[0]}'].fill = SFILL
    mr[0] += 1
def line(key, label, unit, fn, fmt=USD1, bold=False, top=False, inputs=False):
    rr = mr[0]; R[key] = rr
    md[f'A{rr}'] = label; md[f'B{rr}'] = unit
    md[f'A{rr}'].font = BOLD if bold else BLACK
    for i, c in enumerate(C):
        v = fn(i, c, C[i - 1] if i else None)
        md[f'{c}{rr}'] = v
        isf = isinstance(v, str) and v.startswith('=')
        colr = ('008000' if re.fullmatch(r"=Assumptions!\$B\$\d+", v) else '000000') if isf else '0000FF'
        md[f'{c}{rr}'].font = Font(name=F, bold=bold, color=colr)
        md[f'{c}{rr}'].number_format = fmt
        if top: md[f'{c}{rr}'].border = TOP
    mr[0] += 1
X = lambda k: REF[k]
ref = lambda key, c: f'{c}{R[key]}'

hdr('Drivers')
line('idx', 'Quarter number', '', lambda i, c, p: i + 1, '0')
line('days', 'Days', '', lambda i, c, p: [92, 92, 90, 91, 92, 92][i], '0')
line('year', 'Fiscal year', '', lambda i, c, p: '2026' if i < 2 else '2027', '@')
line('rr', 'Reserve return', '%', lambda i, c, p: [f"={X('rr_q3')}", f"={X('rr_q4')}", f"={X('b1')}-{X('gap')}", f"={X('b2')}-{X('gap')}", f"={X('b3')}-{X('gap')}", f"={X('b4')}-{X('gap')}"][i], PCT2)
line('path', 'USDC on the JPMorgan path, before rate drag', '$bn', lambda i, c, p: None if i == 0 else (f"={X('usdc_q4')}" if i == 1 else f"={p}{mr[0]}*(1+{X('g')})^0.25"), BN)
line('usdc', 'USDC average', '$bn', lambda i, c, p: f"={X('usdc_q3')}" if i == 0 else (f"={ref('path', c)}" if i == 1 else f"={ref('path', c)}*(1-{X('drag')}*({ref('rr', c)}-{X('today')})*100)"), BN, bold=True)
line('end', 'USDC at period end', '$bn', lambda i, c, p: f"={X('end_q3')}" if i == 0 else (f"={X('end_q4')}" if i == 1 else f"={ref('usdc', c)}*(1+{X('g')})^0.125"), BN)
line('arc', 'ARC token revenue', '$m', lambda i, c, p: [f"={X('arc_q3')}", f"={X('arc_q4')}", f"={X('arc_27')}/2", f"={X('arc_27')}/2", 0, 0][i])

hdr('Income statement')
line('ri', 'Reserve income', '$m', lambda i, c, p: f"={ref('usdc', c)}*{ref('rr', c)}*1000*{ref('days', c)}/365", bold=True)
line('kept', 'Reserve income Circle keeps', '$m', lambda i, c, p: f"={ref('usdc', c)}*({X('y0')}*{X('k0')}+({ref('rr', c)}-{X('y0')})*{X('mu')})*1000*{ref('days', c)}/365")
line('dtc', 'Distribution and transaction costs', '$m', lambda i, c, p: f"={ref('ri', c)}-{ref('kept', c)}")
line('keep', 'Keep (share of reserve income)', '%', lambda i, c, p: f"=IF({ref('ri', c)}=0,0,{ref('kept', c)}/{ref('ri', c)})", PCT)
line('oth', 'Other revenue ex-ARC', '$m', lambda i, c, p: f"=IF({ref('year', c)}=\"2026\",{X('oth_h2')}/2,{X('oth_h2')}/2*(1+{X('oth_g')}))")
line('arc2', 'ARC token revenue', '$m', lambda i, c, p: f"={ref('arc', c)}")
line('rev', 'Total revenue', '$m', lambda i, c, p: f"={ref('ri', c)}+{ref('oth', c)}+{ref('arc2', c)}", bold=True, top=True)
line('dtc2', 'Less distribution and transaction costs', '$m', lambda i, c, p: f"=-{ref('dtc', c)}")
line('oc', 'Less other costs', '$m', lambda i, c, p: f"=-{X('oc')}")
line('rldc', 'Revenue less distribution costs (RLDC)', '$m', lambda i, c, p: f"={ref('rev', c)}+{ref('dtc2', c)}+{ref('oc', c)}", bold=True, top=True)
line('rldc_m', 'RLDC margin', '%', lambda i, c, p: f"=IF({ref('rev', c)}=0,0,{ref('rldc', c)}/{ref('rev', c)})", PCT)
line('adj', 'Adjusted opex', '$m', lambda i, c, p: f"=IF({ref('year', c)}=\"2026\",({X('opex26')}-{X('opexh1')})/2,{X('opex26')}/4*(1+{X('opex_g')}))")
line('ebitda', 'Adjusted EBITDA', '$m', lambda i, c, p: f"={ref('rldc', c)}-{ref('adj', c)}", bold=True)
line('sbc', 'Stock pay', '$m', lambda i, c, p: f"={X('sbc')}*(1+{X('sbc_g')})^{ref('idx', c)}")
line('da', 'Depreciation and amortisation', '$m', lambda i, c, p: f"={X('da')}")
line('othop', 'Other GAAP opex items', '$m', lambda i, c, p: f"={X('oth_op')}")
line('op', 'Operating income', '$m', lambda i, c, p: f"={ref('ebitda', c)}-{ref('sbc', c)}-{ref('da', c)}-{ref('othop', c)}", bold=True, top=True)
line('cash_open', 'Corporate cash at start of quarter', '$m', lambda i, c, p: f"={X('cash0')}" if i == 0 else f"={p}{{CASH_END}}")
line('int', 'Interest on corporate cash', '$m', lambda i, c, p: f"={ref('cash_open', c)}*({ref('rr', c)}+{X('gap')})*{ref('days', c)}/365")
line('pre', 'Pretax income', '$m', lambda i, c, p: f"={ref('op', c)}+{ref('int', c)}", top=True)
line('taxr', 'Tax rate', '%', lambda i, c, p: f"=IF({ref('year', c)}=\"2026\",{X('t26')},{X('t27')})", PCT)
line('tax', 'Tax', '$m', lambda i, c, p: f"={ref('pre', c)}*{ref('taxr', c)}")
line('ni', 'Net income', '$m', lambda i, c, p: f"={ref('pre', c)}-{ref('tax', c)}", bold=True, top=True)
line('drift', 'Share drift', '%', lambda i, c, p: 0 if i == 0 else f"={X('drift')}", PCT2)
line('binsh', 'Binance shares issued (weighted)', 'm', lambda i, c, p: [f"={X('bin_sh')}*14/92", f"={X('bin_sh')}*78/92", 0, 0, 0, 0][i], '#,##0.000')
line('tazsh', 'Tazapay shares (weighted)', 'm', lambda i, c, p: [0, 0, 0, f"={X('taz_sh')}*0.5", f"={X('taz_sh')}*0.5", 0][i], '#,##0.000')
line('sh', 'Diluted shares', 'm', lambda i, c, p: f"={X('sh0') if i == 0 else p + str(mr[0])}*(1+{ref('drift', c)})+{ref('binsh', c)}+{ref('tazsh', c)}", '#,##0.0')
line('eps', 'GAAP diluted EPS', '$', lambda i, c, p: f"={ref('ni', c)}/{ref('sh', c)}", PS, bold=True)
line('eps_x', 'EPS ex-ARC token revenue', '$', lambda i, c, p: f"=({ref('ni', c)}-{ref('arc2', c)}*(1-{ref('taxr', c)}))/{ref('sh', c)}", PS)

hdr('Cash flow')
line('cf_ni', 'Net income', '$m', lambda i, c, p: f"={ref('ni', c)}")
line('cf_sbc', 'Add stock pay', '$m', lambda i, c, p: f"={ref('sbc', c)}")
line('cf_da', 'Add depreciation and amortisation', '$m', lambda i, c, p: f"={ref('da', c)}")
line('cf_arc', 'Less ARC revenue (cash came in 2026 Q2-Q3)', '$m', lambda i, c, p: f"=-{ref('arc2', c)}")
line('cf_recv', 'ARC receivable collected', '$m', lambda i, c, p: f"={X('recv')}" if i == 0 else 0)
line('cfo', 'Cash from operations', '$m', lambda i, c, p: f"=SUM({c}{R['cf_ni']}:{c}{R['cf_recv']})", bold=True, top=True)
line('cf_capex', 'Capex', '$m', lambda i, c, p: f"=-{X('capex')}")
line('fcf', 'Free cash flow', '$m', lambda i, c, p: f"={ref('cfo', c)}+{ref('cf_capex', c)}", bold=True)
line('cf_bin', 'Binance share placement', '$m', lambda i, c, p: f"={X('bin_cash')}" if i == 0 else 0)
line('cf_opt', 'Option exercise proceeds', '$m', lambda i, c, p: f"={X('opt')}")
line('cf_rsu', 'RSU tax withholding', '$m', lambda i, c, p: f"=-{X('rsu')}")
line('cff', 'Cash from financing (ex USDC flows)', '$m', lambda i, c, p: f"=SUM({c}{R['cf_bin']}:{c}{R['cf_rsu']})", top=True)
line('cash_end', 'Corporate cash at end of quarter', '$m', lambda i, c, p: f"={ref('cash_open', c)}+{ref('fcf', c)}+{ref('cff', c)}", bold=True)
# fill the forward reference to end cash
for i, c in enumerate(C[1:], 1):
    md[f'{c}{R["cash_open"]}'] = f"={C[i - 1]}{R['cash_end']}"

hdr('Balance sheet (period end)')
line('bs_cash', 'Corporate cash', '$m', lambda i, c, p: f"={ref('cash_end', c)}")
line('bs_seg', 'Segregated corporate stablecoins', '$m', lambda i, c, p: f"={X('seg')}")
line('bs_res', 'Reserves for USDC holders', '$m', lambda i, c, p: f"={ref('end', c)}*1000+{X('excess')}")
line('bs_ppe', 'Fixed assets, intangibles, goodwill', '$m', lambda i, c, p: f"={X('ppe0') if i == 0 else p + str(mr[0])}+{X('capex')}-{ref('da', c)}" + (f"+{X('taz')}" if i == 3 else ''))
line('bs_oa', 'Other assets', '$m', lambda i, c, p: f"={X('oa0') if i == 0 else p + str(mr[0])}-{ref('cf_recv', c)}")
line('bs_ta', 'Total assets', '$m', lambda i, c, p: f"=SUM({c}{R['bs_cash']}:{c}{R['bs_oa']})", bold=True, top=True)
line('bs_dep', 'Deposits from USDC holders', '$m', lambda i, c, p: f"={ref('end', c)}*1000")
line('bs_def', 'Deferred ARC presale revenue', '$m', lambda i, c, p: f"={X('def0') if i == 0 else p + str(mr[0])}-{ref('arc2', c)}")
line('bs_ol', 'Other liabilities', '$m', lambda i, c, p: f"={X('ol0')}")
line('bs_tl', 'Total liabilities', '$m', lambda i, c, p: f"=SUM({c}{R['bs_dep']}:{c}{R['bs_ol']})", bold=True, top=True)
line('bs_eq', 'Total equity (roll-forward)', '$m', lambda i, c, p: f"={X('eq0') if i == 0 else p + str(mr[0])}+{ref('ni', c)}+{ref('sbc', c)}+{ref('cf_bin', c)}+{ref('cf_opt', c)}+{ref('cf_rsu', c)}" + (f"+{X('taz')}" if i == 3 else ''), bold=True)
line('check', 'Check: assets less liabilities less equity', '$m', lambda i, c, p: f"=ROUND({ref('bs_ta', c)}-{ref('bs_tl', c)}-{ref('bs_eq', c)},6)", '#,##0.00;(#,##0.00);0.00')
md.column_dimensions['A'].width = 46; md.column_dimensions['B'].width = 7
for c in C: md.column_dimensions[c].width = 12
md.freeze_panes = 'C4'
md[f'A{mr[0] + 1}'] = 'Tazapay closes in 2027; it enters the balance sheet in Q2-27 at its $400m price (paid in shares), with the shares weighted in over Q2-27 and Q3-27.'
md[f'A{mr[0] + 2}'] = 'Everything not shown on the balance sheet is held at 30 June 2026. Stablecoin reserves and the USDC liability move one for one.'
for k in (1, 2): md[f'A{mr[0] + k}'].font = Font(name=F, italic=True, size=9)
tb = mr[0] + 4
md[f'A{tb}'] = 'Tax check (2027 rate from the NOL schedule; feeds nothing, to avoid a circular reference)'; md[f'A{tb}'].font = BOLD
tc = [('FY26 pretax ($m)', f"={X('pre_h1')}+C{R['pre']}+D{R['pre']}"),
      ('NOLs left after 2026 ($m)', f"=MAX(0,{X('nol0')}-0.8*B{tb + 1})"),
      ('FY27 pretax ($m)', f"=SUM(E{R['pre']}:H{R['pre']})"),
      ('NOLs used in 2027 ($m)', f"=MIN(B{tb + 2},0.8*B{tb + 3})"),
      ('2027 rate implied', f"={X('t26')}+0.21*MAX(0,B{tb + 3}-B{tb + 4})/B{tb + 3}"),
      ('2027 rate used (Assumptions)', f"={X('t27')}")]
for j, (lab, f_) in enumerate(tc):
    md[f'A{tb + 1 + j}'] = lab; md[f'B{tb + 1 + j}'] = f_; md[f'B{tb + 1 + j}'].number_format = PCT if 'rate' in lab else USD1

# ---------------- Consensus ----------------
cs = wb.create_sheet('Consensus')
cs['A1'] = 'My forecast against the Street (consensus read 28 Oct 2026; EPS is GAAP diluted)'; cs['A1'].font = Font(name=F, bold=True, size=12)
cs['A3'] = 'H1-26 actuals (Circle 10-Qs)'; cs['A3'].font = BOLD
h1 = [('h1_rev', 'Total revenue', 'total_revenue'), ('h1_ri', 'Reserve income', 'reserve_income'), ('h1_dtc', 'Distribution and transaction costs', 'distribution_transaction_costs'), ('h1_eps', 'GAAP diluted EPS', 'diluted_eps')]
H = {}
for j, (k, lab, col) in enumerate(h1):
    rr = 4 + j; H[k] = f'$B${rr}'
    cs[f'A{rr}'] = lab; cs[f'B{rr}'] = num(ACT['Q1-26'][col]) + num(ACT['Q2-26'][col])
    cs[f'B{rr}'].font = BLUE; cs[f'B{rr}'].number_format = PS if col == 'diluted_eps' else USD1
    cs[f'C{rr}'] = f"Q1-26 {ACT['Q1-26'][col]} + Q2-26 {ACT['Q2-26'][col]}"; cs[f'C{rr}'].font = Font(name=F, size=9, color='555555')
hd = ['Period', 'My revenue ($m)', 'Consensus revenue ($m, Yahoo)', 'Revenue gap', 'My GAAP EPS', 'Consensus EPS (Yahoo)', 'Consensus EPS (Nasdaq/Zacks)', 'EPS gap vs Yahoo', 'My EPS ex-ARC', 'My keep', 'Analysts (Yahoo EPS)']
for j, t in enumerate(hd):
    cl = cs.cell(row=10, column=j + 1, value=t); cl.font = HEAD; cl.fill = HFILL; cl.alignment = Alignment(wrap_text=True, horizontal='center')
def M(key, col): return f"Model!{col}{R[key]}"
rowsdef = {
    'Q3-26': dict(rev=f"={M('rev', 'C')}", eps=f"={M('eps', 'C')}", epsx=f"={M('eps_x', 'C')}", keep=f"={M('keep', 'C')}"),
    'Q4-26': dict(rev=f"={M('rev', 'D')}", eps=f"={M('eps', 'D')}", epsx=f"={M('eps_x', 'D')}", keep=f"={M('keep', 'D')}"),
    'FY26': dict(rev=f"={H['h1_rev']}+{M('rev', 'C')}+{M('rev', 'D')}", eps=f"={H['h1_eps']}+{M('eps', 'C')}+{M('eps', 'D')}",
                 epsx=f"={H['h1_eps']}+{M('eps_x', 'C')}+{M('eps_x', 'D')}",
                 keep=f"=1-({H['h1_dtc']}+{M('dtc', 'C')}+{M('dtc', 'D')})/({H['h1_ri']}+{M('ri', 'C')}+{M('ri', 'D')})"),
    'FY27': dict(rev=f"=SUM(Model!E{R['rev']}:H{R['rev']})", eps=f"=SUM(Model!E{R['eps']}:H{R['eps']})", epsx=f"=SUM(Model!E{R['eps_x']}:H{R['eps_x']})",
                 keep=f"=1-SUM(Model!E{R['dtc']}:H{R['dtc']})/SUM(Model!E{R['ri']}:H{R['ri']})"),
}
for j, (p, d) in enumerate(rowsdef.items()):
    rr = 11 + j; c = CONS[p]
    vals = [p, d['rev'], num(c['revenue_musd_yahoo']), f"=B{rr}/C{rr}-1", d['eps'], num(c['eps_yahoo']), num(c['eps_nasdaq']), f"=E{rr}/F{rr}-1", d['epsx'], d['keep'], int(c['analysts_eps_yahoo'])]
    fmts = ['@', USD1, USD1, PCT, PS, '$0.000', PS, PCT, PS, PCT, '0']
    for k, (v, f_) in enumerate(zip(vals, fmts)):
        cl = cs.cell(row=rr, column=k + 1, value=v); cl.number_format = f_
        cl.font = (GREEN if isinstance(v, str) and v.startswith('=Model!') else BLACK) if isinstance(v, str) and v.startswith('=') else (BLUE if k > 0 else BOLD)
cs['A16'] = 'Sources: Yahoo Finance earningsTrend (query2.finance.yahoo.com); Nasdaq analyst earnings forecast (Zacks data), api.nasdaq.com. Both read 28 Oct 2026.'
cs['A17'] = f"Q4-26: my EPS is above the Street on about the same revenue. On the Street's revenue the gap needs ~$32m more Q4 cost: a keep near 31% (mine 35.7%), adjusted opex near $184m (guide leaves ~$152m), or less interest on Circle's own cash. See model.txt."
cs['A18'] = f"FY27 tax: {_M.T27:.1f}% from the NOL schedule ($478.8m federal NOLs at end-2025, 80% limit), not the 21% statutory rate. A valuation-allowance release would be a one-off benefit and is left out."
for k in (16, 17, 18): cs[f'A{k}'].font = Font(name=F, italic=True, size=9)
cs.column_dimensions['A'].width = 34
for col in 'BCDEFGHIJK': cs.column_dimensions[col].width = 15
cs.row_dimensions[10].height = 45

# ---------------- Valuation ----------------
vs = wb.create_sheet('Valuation')
vs['A1'] = 'Price target: scenario-weighted value per share, rolled forward one year at the cost of equity'; vs['A1'].font = Font(name=F, bold=True, size=12)
vs['A2'] = 'Each scenario value is a 10-year cash-flow value from scripts/valuation.py (USDC path from each forecast, forward-curve yields, measured keep, the add-ons below). Changing the cost of equity here moves the roll-forward only; re-run valuation.py to move the scenario values.'
vs['A2'].font = Font(name=F, italic=True, size=9); vs['A2'].alignment = Alignment(wrap_text=True); vs.merge_cells('A2:E2'); vs.row_dimensions[2].height = 42
for j, t in enumerate(['Scenario', 'Stablecoin market', 'Value per share', 'My weight (set 27 Oct)', 'Options-market odds (cross-check)']):
    cl = vs.cell(row=4, column=j + 1, value=t); cl.font = HEAD; cl.fill = HFILL
for j, k in enumerate(('flat', 'jpmorgan', 'standard_chartered')):
    rr = 5 + j
    vs[f'A{rr}'] = k; vs[f'B{rr}'] = VAL[k]['Stablecoin market']
    vs[f'C{rr}'] = num(VAL[k]['Value per share ($)']); vs[f'C{rr}'].font = BLUE; vs[f'C{rr}'].number_format = PS
    vs[f'D{rr}'] = VS['w_' + k]; vs[f'D{rr}'].font = BLUE; vs[f'D{rr}'].fill = YELLOW; vs[f'D{rr}'].number_format = '0%'
    vs[f'E{rr}'] = round(VS['wo_' + k], 4); vs[f'E{rr}'].font = BLUE; vs[f'E{rr}'].number_format = '0.0%'
rows = [
    ('Weighted value today', '=SUMPRODUCT(C5:C7,D5:D7)', PS),
    ('Weights add to', '=SUM(D5:D7)', '0%'),
    ('Risk-free rate (10-year Treasury)', CD['risk_free_10y_pct'] / 100, PCT2),
    ('Equity risk premium (Damodaran)', CD['equity_risk_premium_pct'] / 100, PCT2),
    ('CRCL weekly beta vs S&P 500, 50 weeks', COE['beta_weekly'], '0.000'),
    ('Blume-adjusted beta', '=0.67*C13+0.33', '0.000'),
    ('Cost of equity (CAPM; Circle has no debt)', '=C11+C14*C12', PCT),
    ('12-month price target', '=C9*(1+C15)', PS),
    ('CRCL price', f"={X('price')}", PS),
    ('Target vs price', '=C16/C17-1', PCT),
    ('Weight on Standard Chartered the price implies (rest at my flat:JPMorgan ratio)', '=(C17-(D5*C5+D6*C6)/(D5+D6))/(C7-(D5*C5+D6*C6)/(D5+D6))', PCT),
    ('Street mean target (Yahoo, 24 analysts)', m['street_mean_target'], PS),
    ('Nasdaq consensus target', 109.56, PS),
    ('Target / my 2027 GAAP EPS (x)', '=C16/Consensus!E14', '0.0x'),
    ('Target / my 2027 EPS ex-ARC (x)', '=C16/Consensus!I14', '0.0x'),
    ('Rating range: HOLD while the price is between', '=C16/1.15', PS),
    ('... and', '=C16/0.85', PS),
]
vs['E9'] = '=SUMPRODUCT(C5:C7,E5:E7)/SUM(E5:E7)'; vs['E9'].number_format = PS; vs['E16'] = '=E9*(1+C15)'; vs['E16'].number_format = PS; vs['E18'] = '=E16/C17-1'; vs['E18'].number_format = PCT
for j, (lab, v, f_) in enumerate(rows):
    rr = 9 + j
    vs[f'A{rr}'] = lab; vs[f'C{rr}'] = v; vs[f'C{rr}'].number_format = f_
    vs[f'C{rr}'].font = (GREEN if 'Assumptions!' in v else BLACK) if isinstance(v, str) and v.startswith('=') else BLUE
for k in (9, 16): vs[f'A{k}'].font = BOLD; vs[f'C{k}'].font = Font(name=F, bold=True)
vs['D11'] = 'US Treasury par curve, 27 Oct 2026'; vs['D12'] = 'implied ERP, 1 Oct 2026, pages.stern.nyu.edu/~adamodar'; vs['D13'] = 'beta.py on Yahoo weekly closes'
vs['D21'] = 'Nasdaq, read 28 Oct 2026 (12 buy, 6 hold, 3 sell)'
for k in (11, 12, 13, 21): vs[f'D{k}'].font = Font(name=F, size=9, color='555555')
vs['A27'] = 'Standard Chartered weight vs value (rest at my 35:55 flat:JPMorgan ratio)'; vs['A27'].font = BOLD; vs['A27'].fill = SFILL
for j, t in enumerate(['Weight on Standard Chartered', 'Value today', '12-month target', 'vs price']):
    cl = vs.cell(row=28, column=j + 1, value=t); cl.font = HEAD; cl.fill = HFILL
for j, w in enumerate((0.0, 0.05, '=D7', '=E7', 0.20)):
    rr = 29 + j
    vs[f'A{rr}'] = w; vs[f'A{rr}'].font = BLUE; vs[f'A{rr}'].number_format = '0%'
    vs[f'B{rr}'] = f'=(1-A{rr})*(D5*C5+D6*C6)/(D5+D6)+A{rr}*C7'; vs[f'B{rr}'].number_format = PS
    vs[f'C{rr}'] = f'=B{rr}*(1+$C$15)'; vs[f'C{rr}'].number_format = PS
    vs[f'D{rr}'] = f'=C{rr}/$C$17-1'; vs[f'D{rr}'].number_format = PCT
vs['A35'] = f"Target uses my weights, set on 27 Oct and unchanged. Cross-check: the Jan-2028 options' odds of CRCL finishing below ${VS['band_low']:.0f}, between, and above ${VS['band_high']:.0f} (scenario values rolled forward at the cost of equity, split at geometric midpoints). Options odds are risk-neutral and lean to the downside."
vs['A36'] = f"Added to every scenario once: cash and receivables $1,850m, 25% of Arc at the $3bn presale valuation ($750m), Tazapay at $400m, federal NOL shield ${VS['nol_shield_musd']:.0f}m; 274.9m shares."
for k in (35, 36): vs[f'A{k}'].font = Font(name=F, italic=True, size=9)
vs['A38'] = 'USDC share sensitivity (every scenario holds USDC at 23.7% of stablecoins; valuation.py)'; vs['A38'].font = BOLD; vs['A38'].fill = SFILL
for j, (lab, key) in enumerate((('USDC loses 0.5 point of share a year: value today', 'share_loss_05_value'), ('USDC loses 1.0 point of share a year: value today', 'share_loss_10_value'))):
    rr = 39 + j
    vs[f'A{rr}'] = lab; vs[f'C{rr}'] = round(VS[key], 2); vs[f'C{rr}'].font = BLUE; vs[f'C{rr}'].number_format = PS
    vs[f'D{rr}'] = f'=C{rr}*(1+$C$15)'; vs[f'D{rr}'].number_format = PS
vs['D38'] = '12-month target'
vs.column_dimensions['A'].width = 62; vs.column_dimensions['B'].width = 34; vs.column_dimensions['C'].width = 14; vs.column_dimensions['D'].width = 40

# ---------------- Hyperliquid ----------------
hs = wb.create_sheet('Hyperliquid')
hs['A1'] = "Hyperliquid's USDC, measured on HyperEVM, and what it does to Circle's Q3-26 keep"; hs['A1'].font = Font(name=F, bold=True, size=12)
hs['A2'] = 'Daily balances of Coinbase\'s AQA treasury (0xc20699185c15D0a2fD65779BB5d69f5b0B113c00) and Circle\'s Core Deposit Wallet (0x6b9e773128f453f5c2c60935ee2de2cbc5390a24), archive RPC, data/hyperliquid_split_onchain_daily.csv. Arbitrum bridge from DefiLlama.'
hs['A2'].font = Font(name=F, italic=True, size=9)
hrows = [
    ('Coinbase treasury, Q2-26 average ($bn)', HL['coinbase_treasury_q2_bn'], BN),         # 4
    ("Circle's wallet, Q2-26 average ($bn)", HL['circle_wallet_q2_bn'], BN),                # 5
    ('Legacy Arbitrum bridge, Q2-26 average ($bn)', HL['arbitrum_bridge_q2_bn'], BN),       # 6
    ('Coinbase treasury, Q3-26 average ($bn)', HL['coinbase_treasury_q3_bn'], BN),          # 7
    ("Circle's wallet, Q3-26 average ($bn)", HL['circle_wallet_q3_bn'], BN),                # 8
    ('Legacy Arbitrum bridge, Q3-26 average ($bn)', HL['arbitrum_bridge_q3_bn'], BN),       # 9
    ('Reserve return, Q3-26', HL['q3_return_pct'] / 100, PCT2),                              # 10
    ("Circle's take on a shared-pool dollar, low", HL['shared_take_low_pct'] / 100, PCT2),   # 11
    ("Circle's take on a shared-pool dollar, high", HL['shared_take_high_pct'] / 100, PCT2), # 12
    ("Coinbase's take of the shared pool", HL['coinbase_take_of_shared_pct'] / 100, PCT),    # 13
    ('$m per $bn-year for 92 days', '=92/365*1000', '0.000'),                                # 14
    ('Circle keeps on Hyperliquid USDC, Q2 mix, low take ($m)', '=(B5*B10+B6*B11)*B14', USD1),   # 15
    ('Circle keeps on Hyperliquid USDC, Q3 mix, low take ($m)', '=(B8*B10+B9*B11)*B14', USD1),   # 16
    ('Change, low take ($m)', '=B16-B15', USD1),                                              # 17
    ('Change, high take ($m)', '=(B8*B10+B9*B12)*B14-(B5*B10+B6*B12)*B14', USD1),            # 18
    ('Circle pays Coinbase on Hyperliquid USDC, change Q2 to Q3 mix ($m)', '=((B7*B10+B9*B13*B10)-(B4*B10+B6*B13*B10))*B14', USD1),  # 19
    ("Binance customers' USDC, Q3 vs Q2 average ($bn)", HL['binance_customer_usdc_change_bn'], BN),  # 20
    ('Circle keeps on that USDC, low end (1.34c a year), $m', '=B20*0.0134*B14', USD1),        # 21
    ('Keep, Q2-26 (Circle 10-Q)', HL['keep_q2_pct'] / 100, PCT),                              # 22
    ('Pre-registered Q3 reserve income ($m)', HL['q3_reserve_income_preregistered_musd'], USD),  # 23
    ('Keep, Q3-26, used in the model (midpoint plus the backtest correction)', '=B27+B28', PCT),  # 24
    ('Keep, Q3-26, method low (keep_method.py: Binance counted)', HL['q3_keep_low'] / 100, PCT),  # 25
    ('Keep, Q3-26, method high (keep_method.py: Binance not counted)', HL['q3_keep_high'] / 100, PCT),  # 26
    ('Keep, Q3-26, method midpoint', '=(B25+B26)/2', PCT),                                    # 27
    ('Backtest correction: the method ran this far low on average', -BT['method_bias_pt'] / 100, PCT2),  # 28
]
for j, (lab, v, f_) in enumerate(hrows):
    rr = 4 + j
    hs[f'A{rr}'] = lab; hs[f'B{rr}'] = v; hs[f'B{rr}'].number_format = f_
    hs[f'B{rr}'].font = BLACK if isinstance(v, str) else BLUE
for k in (24, 25, 26): hs[f'A{k}'].font = BOLD
hs['A30'] = "Before 13 June 2026 all native Hyperliquid USDC sat in Circle's Core Deposit Wallet (Circle's platform). Since then the two addresses are held at 90/10."
hs['A30'].font = Font(name=F, italic=True, size=9)
hs.column_dimensions['A'].width = 66; hs.column_dimensions['B'].width = 14

# ---------------- Options ----------------
op = wb.create_sheet('Options')
op['A1'] = 'CRCL options, Cboe delayed quotes (data/crcl_options_20261028.csv): what the market pays for'; op['A1'].font = Font(name=F, bold=True, size=12)
for j, t in enumerate(['Expiry', 'Measure', 'Strike or level ($)', 'Probability']):
    cl = op.cell(row=3, column=j + 1, value=t); cl.font = HEAD; cl.fill = HFILL
txt = open(os.path.join(RESULTS, 'options_check.txt')).read().splitlines()
rr = 4; exp = ''
for ln in txt:
    mm = re.match(r'(\d{8}):', ln)
    if mm: exp = f'{mm.group(1)[:4]}-{mm.group(1)[4:6]}-{mm.group(1)[6:]}'; continue
    mm = re.search(r'call spread (\d+)-(\d+): odds of finishing above ~(\d+) = ([\d.]+)%', ln)
    if mm:
        op.append([exp, f'call spread {mm.group(1)}-{mm.group(2)}', float(mm.group(3)), float(mm.group(4)) / 100]); rr += 1; continue
    mm = re.search(r'lognormal at (\d+)%: P\(above \$(\d+)\) = ([\d.]+)%', ln)
    if mm:
        op.append([exp, f'lognormal at {mm.group(1)}% implied vol', float(mm.group(2)), float(mm.group(3)) / 100]); rr += 1
for row in op.iter_rows(min_row=4, max_row=op.max_row):
    for cl in row: cl.font = BLUE
    row[3].number_format = PCT; row[2].number_format = '$#,##0'
st = [l for l in txt if 'straddle' in l]
if st: op.cell(row=op.max_row + 2, column=1, value=st[0].strip()).font = Font(name=F, italic=True, size=9)
op.column_dimensions['A'].width = 14; op.column_dimensions['B'].width = 30; op.column_dimensions['C'].width = 18; op.column_dimensions['D'].width = 12

# ---------------- Line items ----------------
li = wb.create_sheet('Line items', 4)
li['A1'] = 'Line by line: my forecast, Circle FY26 guidance (Q2-26 call), and the Street where free sources carry it (revenue, EPS)'; li['A1'].font = Font(name=F, bold=True, size=12)
LIR = list(__import__('csv').reader(open(os.path.join(RESULTS, 'table_line_items.csv'))))
for i, row in enumerate(LIR):
    for j, v in enumerate(row):
        cl = li.cell(row=3 + i, column=j + 1, value=v)
        cl.font = HEAD if i == 0 else (BOLD if j == 0 else BLUE)
        if i == 0: cl.fill = HFILL; cl.alignment = Alignment(wrap_text=True, horizontal='center')
li.column_dimensions['A'].width = 34
for col in 'BCDEFGHIJ': li.column_dimensions[col].width = 15
li.column_dimensions['E'].width = 26
li.row_dimensions[3].height = 32
li.cell(row=4 + len(LIR), column=1, value='Consensus adjusted EBITDA, RLDC and distribution costs are not on free sources; the guidance column and the Q4 walk (model.txt) stand in for them.').font = Font(name=F, italic=True, size=9)

# ---------------- Backtest ----------------
bt = wb.create_sheet('Backtest')
bt['A1'] = 'Does the keep method work? Each quarter predicted from the quarter before (backtest_keep.py)'; bt['A1'].font = Font(name=F, bold=True, size=12)
hd = ['Quarter', 'Actual keep', 'Method low', 'Method high', 'Method midpoint', 'Same as last quarter', 'Method + rate table', 'Rate table only', 'Method miss (pt)', 'Naive miss (pt)']
for j, t in enumerate(hd):
    cl = bt.cell(row=3, column=j + 1, value=t); cl.font = HEAD; cl.fill = HFILL; cl.alignment = Alignment(wrap_text=True, horizontal='center')
for i, r in enumerate(BTQ):
    rr = 4 + i
    bt.cell(row=rr, column=1, value=r['quarter']).font = BOLD
    for j, k in enumerate(('actual_keep_pct', 'method_low_pct', 'method_high_pct', 'method_mid_pct', 'same_as_last_pct', 'method_with_yield_adj_pct', 'yield_adj_only_pct')):
        cl = bt.cell(row=rr, column=j + 2, value=num(r[k]) / 100); cl.font = BLUE; cl.number_format = PCT
    bt.cell(row=rr, column=9, value=f'=(E{rr}-B{rr})*100').number_format = '+0.0;-0.0;0.0'
    bt.cell(row=rr, column=10, value=f'=(F{rr}-B{rr})*100').number_format = '+0.0;-0.0;0.0'
n = len(BTQ); last = 3 + n
bt.cell(row=last + 2, column=1, value='Mean miss (pt)').font = BOLD
for col, letter in ((9, 'I'), (10, 'J')):
    bt.cell(row=last + 2, column=col, value=f'=SUMPRODUCT(ABS({letter}4:{letter}{last}))/{n}').number_format = '0.00'
bt.cell(row=last + 3, column=1, value='Average miss, method (pt)').font = BOLD
bt.cell(row=last + 3, column=9, value=f'=AVERAGE(I4:I{last})').number_format = '+0.00;-0.00'
bt.cell(row=last + 5, column=1, value='Measured flows beat "same as last quarter". The method runs low (Circle\'s own platform kept growing in ways the chain does not show), so the Q3-26 keep adds back the average miss. Circle\'s rate table made the past four quarters worse.').font = Font(name=F, italic=True, size=9)
bt.column_dimensions['A'].width = 26
for col in 'BCDEFGHIJ': bt.column_dimensions[col].width = 13
bt.row_dimensions[3].height = 45

# ---------------- Drag ----------------
dg = wb.create_sheet('Drag')
dg['A1'] = 'How much float higher bill rates cost (rate_drag.py)'; dg['A1'].font = Font(name=F, bold=True, size=12)
dl = open(os.path.join(RESULTS, 'rate_drag.txt')).read().splitlines()
for i, ln in enumerate(dl): dg[f'A{3 + i}'] = ln; dg[f'A{3 + i}'].font = Font(name=F)
dg.column_dimensions['A'].width = 140

# ---------------- Checks ----------------
ck = wb.create_sheet('Checks', 0)
ck['A1'] = 'Circle (CRCL) model: read me first'; ck['A1'].font = Font(name=F, bold=True, size=13)
notes = ['Sheets: Assumptions (every input with its source), Model (income statement, cash flow, balance sheet, Q3-26 to Q4-27, tax check), Consensus, Line items, Valuation, Hyperliquid, Backtest, Drag, Options.',
         'Colours: blue input, black formula, green link to another sheet, yellow key judgement.',
         'Built by scripts/build_workbook.py from the same data as scripts/model.py. The two should agree to the cent.']
for j, t in enumerate(notes): ck[f'A{3 + j}'] = t; ck[f'A{3 + j}'].font = Font(name=F)
chk = [('Largest balance sheet gap, any quarter ($m)', f"=MAX(ABS(MIN(Model!C{R['check']}:H{R['check']})),ABS(MAX(Model!C{R['check']}:H{R['check']})))", '0.000000'),
       ('Q3-26 keep used (measured, backtest-corrected)', '=Hyperliquid!B24', PCT),
       ('Q3-26 GAAP EPS', f"=Model!C{R['eps']}", PS),
       ('FY27 GAAP EPS', '=Consensus!E14', PS),
       ('12-month price target', '=Valuation!C16', PS),
       ('Target vs price', '=Valuation!C18', PCT)]
for j, (lab, f_, fm) in enumerate(chk):
    rr = 8 + j; ck[f'A{rr}'] = lab; ck[f'B{rr}'] = f_; ck[f'B{rr}'].font = GREEN; ck[f'B{rr}'].number_format = fm
ck.column_dimensions['A'].width = 60; ck.column_dimensions['B'].width = 14

out = os.path.join(RESULTS, 'circle_model.xlsx')
wb.save(out)
print('wrote', out)
